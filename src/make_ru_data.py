"""Extractor for the Russian marketplaces. Reads the official files saved on 2026-10-26 in /workspace/ru-raw
(not committed: 20+ MB) and writes compact rate files under src/data/{ozon,wildberries,yandex-market}/ru.json.
  Ozon:  'Таблица категорий для расчёта вознаграждения' XLSX (in force from 28 Aug 2026; the 7 Oct 2026 table only adds FBO PVZ)
  WB:    'Таблица со значениями для расчета КВВ' PDF, in force from 7 Jul 2026 (parsed with pdfplumber into wb_rows.json)
  YM:    'Размещение товаров на витрине' help page, all category tables (parsed into ym_rows.json)
File format: {"groups": [...], "o": [[groupIndex, name, rates], ...]} -> expanded by the engine into {n, g, r}."""
import json, os, re, collections
RAW = os.environ.get('RU_RAW', '/workspace/ru-raw')
D = os.path.join(os.path.dirname(__file__), 'data')
def dump(p, groups, rows):
    gi = {g: i for i, g in enumerate(groups)}
    json.dump({'groups': groups, 'o': [[gi[g], n, r] for g, n, r in rows]}, open(os.path.join(D, p), 'w'), ensure_ascii=False, separators=(',', ':'))
    print(p, len(rows))
P = lambda x: round(x * 100, 2) if isinstance(x, (int, float)) else None

# ---------------- Ozon: one option per category; product types whose rates differ from their category get their own option
def ozon():
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(RAW, 'oz_28082026.xlsx'), read_only=True).worksheets[0]
    rows = [r for i, r in enumerate(ws.iter_rows(values_only=True)) if i >= 2 and r[0]]
    def rates(r):
        v = [P(x) for x in r[3:13]]
        return {'fbo': v[0:3], 'fbs': v[6:9], 'rfbs': v[9]}
    by = collections.defaultdict(list)
    for r in rows: by[(r[0], r[1])].append(r)
    out = []
    for (main, cat), lst in by.items():
        cnt = collections.Counter(json.dumps(rates(r)) for r in lst)
        top = json.loads(cnt.most_common(1)[0][0])
        out.append((main, cat, top))
        for r in lst:
            rr = rates(r)
            if rr != top: out.append((main, f'{cat} › {r[2]}', rr))
    out.sort(key=lambda x: (x[0], x[1]))
    groups = sorted(set(x[0] for x in out))
    dump('ozon/ru.json', groups, out)
    # FBO/FBS logistics (in force from 30 Oct 2026): volume bucket x dispatch cluster x destination cluster x price (<=300 / >300 RUB)
    wb = openpyxl.load_workbook(os.path.join(RAW, 'oz_logistics_30102026.xlsx'), read_only=True)
    rows = [r for i, r in enumerate(wb.worksheets[0].iter_rows(values_only=True)) if i >= 3 and r[1]]
    dflt = [r for i, r in enumerate(wb.worksheets[1].iter_rows(values_only=True)) if i >= 3 and r[1]]
    buckets = [r[1] for r in dflt]
    def upper(b):
        m = re.findall(r'[\d,]+', b); return 1e9 if b.startswith('От') else float(m[-1].replace(',', '.'))
    clusters = sorted(set(r[2] for r in rows) | set(r[3] for r in rows))
    ci = {c: i for i, c in enumerate(clusters)}; bi = {b: i for i, b in enumerate(buckets)}
    routes = {}
    for r in rows:
        k = f'{ci[r[2]]}-{ci[r[3]]}'
        routes.setdefault(k, [None] * len(buckets))[bi[r[1]]] = [r[4], r[5]]
    d = json.load(open(os.path.join(D, 'ozon/ru.json')))
    d['tables'] = {'logistics': {'upper': [upper(b) for b in buckets], 'labels': buckets, 'clusters': clusters,
                                 'default': [[r[2], r[3]] for r in dflt], 'routes': routes, 'bands': [300]}}
    json.dump(d, open(os.path.join(D, 'ozon/ru.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
    print('ozon logistics routes', len(routes))

# ---------------- Wildberries: one option per subject (предмет)
def wb():
    rows = json.load(open(os.path.join(RAW, 'wb_rows.json')))
    out = sorted(((p, s, {'fbw': v[0], 'fbs': v[1], 'dbs': v[2]}) for p, s, v in rows), key=lambda x: (x[0], x[1]))
    dump('wildberries/ru.json', sorted(set(x[0] for x in out)), out)

# ---------------- Yandex Market: subcategories folded into the nearest listed ancestor when the rates are identical
def ym():
    rows = json.load(open(os.path.join(RAW, 'ym_rows.json')))
    f = lambda s: float(s.rstrip('%'))
    d = {}
    for _, x in rows:
        v = [f(x[2]), f(x[3]), f(x[4]), f(x[5])]
        d[(x[0], x[1])] = {'_': v[0]} if len(set(v)) == 1 else {'fby': v[0], 'fbs': v[1], 'exp': v[2], 'dbs': v[3]}
    out = []
    for (p, c), v in d.items():
        if p == 'Все товары': continue  # page row with the 42% low-price tariff, handled by the low-price rule
        if c and c.count('-') >= 2:  # keep the first two levels; fold deeper levels that repeat their ancestor's rates
            parts = c.split('-'); anc = None
            for k in range(len(parts) - 1, -1, -1):
                key = (p, '-'.join(parts[:k]))
                if key in d: anc = d[key]; break
            if anc == v: continue
        name = re.sub(r'-(?=[А-ЯЁA-Z0-9])', ' › ', c) if c else 'Все товары категории'
        out.append((p, name, v))
    out.sort(key=lambda x: (x[0], '' if x[1] == 'Все товары категории' else x[1]))
    dump('yandex-market/ru.json', sorted(set(x[0] for x in out)), out)

if __name__ == '__main__':
    ozon(); wb(); ym()
